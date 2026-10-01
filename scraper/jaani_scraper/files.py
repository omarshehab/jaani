"""Files linked from an RTI page (source S3, Template D; sections 6.7, 6A.5).

PDF: characters are grouped into runs by font, each table cell / line run is classified and recovered on its own
(never two cells merged before recovery), then the recovered content is laid out as simple HTML and parsed by the
same parser as web pages. Page number and bbox of every cell/line are kept for review crops.
"""
import html as htmlmod
import io
import re
from dataclasses import dataclass, field

import pdfplumber

from . import integrity as ti
from .bijoy import bijoy_to_unicode
from .parse import parse_page
from .textnorm import norm

FILE_CAP = 15 * 1024 * 1024   # linked files are downloaded only up to 15 MB (section 3)


@dataclass
class Region:
    text: str
    page: int
    bbox: tuple
    recovered: str = ''          # '' | 'bijoy'
    unresolved: bool = False
    log: list = field(default_factory=list)


def recover_run(text, font):
    """-> (text, recovered, unresolved, log entry)"""
    cls = ti.classify_text(text, font)
    entry = dict(cls=cls, font=font or '', input=text[:200])
    if cls in (ti.ENGLISH_OK, ti.UNICODE_OK, ti.EMPTY):
        return text, '', False, dict(entry, method='none', decision='as_is')
    if cls == ti.BIJOY_ANSI:
        out = bijoy_to_unicode(text, digits_bengali=True)
        ok = ti.classify_text(out) in ti.STORABLE
        return (out if ok else text), ('bijoy' if ok else ''), not ok, dict(
            entry, method='bijoy_to_unicode', candidate=out[:200], decision='converted' if ok else 'failed_after_conversion')
    return text, '', True, dict(entry, method='none', decision='unresolved')


def _runs(words):
    """Words (one line) -> runs of consecutive words with the same legacy/non-legacy font."""
    runs = []
    for w in words:
        legacy = bool(ti.is_legacy_font(w.get('fontname')))
        if runs and runs[-1][0] == legacy and runs[-1][1] == w.get('fontname'):
            runs[-1][2].append(w['text'])
        else:
            runs.append([legacy, w.get('fontname'), [w['text']]])
    return [(font, ' '.join(t)) for _, font, t in runs]


def _lines(words, tol=3):
    lines = []
    for w in sorted(words, key=lambda w: (round(w['top']), w['x0'])):
        if lines and abs(lines[-1][0] - w['top']) <= tol:
            lines[-1][1].append(w)
        else:
            lines.append([w['top'], [w]])
    return [sorted(ws, key=lambda w: w['x0']) for _, ws in lines]


def _region(words, page_no, bbox):
    texts, rec, unres, logs = [], '', False, []
    for line in _lines(words):
        for font, t in _runs(line):
            out, r, u, entry = recover_run(t, font)
            texts.append(out)
            rec = rec or r
            unres = unres or u
            logs.append(dict(entry, page=page_no, bbox=[round(x, 1) for x in bbox]))
    return Region(norm(' '.join(texts)), page_no, tuple(round(x, 1) for x in bbox), rec, unres, logs)


def _inside(w, bbox):
    cx, cy = (w['x0'] + w['x1']) / 2, (w['top'] + w['bottom']) / 2
    return bbox[0] <= cx <= bbox[2] and bbox[1] <= cy <= bbox[3]


BASE14_RE = re.compile(r'^(?:[A-Z]{6}\+)?(Helvetica|Times|Courier|Arial|GlyphLessFont)', re.I)


def is_scanned_page(page):
    """One image covers >= 85% of the page and every text run is a standard (non-legacy, non-Bengali) font:
    the text layer is a scanner's OCR guess, not the document's text."""
    area = float(page.width * page.height) or 1.0
    big = any((im['x1'] - im['x0']) * (im['bottom'] - im['top']) >= 0.85 * area for im in page.images)
    if not big:
        return False
    fonts = {c.get('fontname', '') for c in page.chars}
    bengali = any('\u0980' <= c.get('text', ' ')[:1] <= '\u09ff' for c in page.chars)
    return not bengali and all(BASE14_RE.match(f or '') for f in fonts)


def pdf_regions(data):
    """-> (items, info). items: [('table', page, top, rows[[Region|None]]) | ('line', page, top, Region)]"""
    items, info = [], dict(pages=0, no_text_pages=[], scanned_pages=[], legacy_runs=0, unresolved_runs=0, recovery=[])
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        info['pages'] = len(pdf.pages)
        for pno, page in enumerate(pdf.pages, 1):
            if not page.chars:
                info['no_text_pages'].append(pno)
                continue
            if is_scanned_page(page):
                # A scanner's own OCR layer (Latin, standard font) over a page image: it is not Bijoy and not
                # the text; the page is read only by 6A.2-C (render + OCR/vision), never from this layer.
                info['scanned_pages'].append(pno)
                info['recovery'].append(dict(page=pno, method='scanned_page', decision='text_layer_ignored',
                                             fonts=sorted({c.get('fontname', '') for c in page.chars})[:5]))
                continue
            words = page.extract_words(extra_attrs=['fontname'], keep_blank_chars=False, use_text_flow=False)
            used = set()
            for t in page.find_tables():
                rows = []
                for r in t.rows:
                    row = []
                    for cell in r.cells:
                        if cell is None:
                            row.append(None)
                            continue
                        ws = [w for i, w in enumerate(words) if _inside(w, cell)]
                        used.update(i for i, w in enumerate(words) if _inside(w, cell))
                        row.append(_region(ws, pno, cell))
                    rows.append(row)
                items.append(('table', pno, t.bbox[1], rows))
            rest = [w for i, w in enumerate(words) if i not in used]
            for line in _lines(rest):
                bbox = (min(w['x0'] for w in line), min(w['top'] for w in line),
                        max(w['x1'] for w in line), max(w['bottom'] for w in line))
                items.append(('line', pno, bbox[1], _region(line, pno, bbox)))
    items.sort(key=lambda it: (it[1], it[2]))
    for it in items:
        regs = [c for row in it[3] for c in row if c] if it[0] == 'table' else [it[3]]
        for rg in regs:
            info['recovery'].extend(rg.log)
            info['legacy_runs'] += sum(1 for e in rg.log if e['method'] == 'bijoy_to_unicode')
            info['unresolved_runs'] += sum(1 for e in rg.log if e['decision'] != 'as_is' and e['decision'] != 'converted')
    return items, info


def _cell_html(rg, tag):
    if rg is None:
        return f'<{tag}></{tag}>'
    attrs = f' data-jaani-rec-text="{rg.recovered}"' if rg.recovered else ''
    return f'<{tag}{attrs}>{htmlmod.escape(rg.text)}</{tag}>'


def regions_to_html(items):
    parts = []
    for it in items:
        if it[0] == 'table':
            parts.append('<table>' + ''.join('<tr>' + ''.join(_cell_html(c, 'td') for c in row) + '</tr>'
                                             for row in it[3]) + '</table>')
        else:
            parts.append(_cell_html(it[3], 'p'))
    return '<html><body>' + '\n'.join(parts) + '</body></html>'


def parse_pdf(data, url):
    """-> (ParsedPage, meta, info). meta['where'][(role, field)] = {file, page, bbox} for review crops."""
    items, info = pdf_regions(data)
    page = parse_page(regions_to_html(items), url)
    page.recovery.extend(info['recovery'])
    if info['legacy_runs']:
        page.flags.append('legacy_font_file')
    if info['unresolved_runs']:
        page.flags.append('bijoy_garbled')
    if info['no_text_pages']:
        page.flags.append('no_text_layer')
    if info['scanned_pages']:
        page.flags.append('scanned_file')
    regions = []
    for it in items:
        regions.extend([c for row in it[3] for c in row if c] if it[0] == 'table' else [it[3]])
    where = {}
    for role, block in page.roles.items():
        for f, fr in block.fields.items():
            if not fr.value:
                continue
            hit = next((rg for rg in regions if fr.value == rg.text), None) or \
                next((rg for rg in regions if fr.value in rg.text), None)
            if hit:
                where[(role, f)] = dict(file=url, page=hit.page, bbox=hit.bbox)
                if hit.recovered:
                    fr.recovered = hit.recovered
    return page, dict(where=where), info


def parse_docx(data, url):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = []

    def para_html(p):
        texts, rec = [], ''
        for run in p.runs:
            out, r, _, _ = recover_run(run.text, run.font.name)
            texts.append(out)
            rec = rec or r
        attrs = f' data-jaani-rec-text="{rec}"' if rec else ''
        return f'<p{attrs}>{htmlmod.escape("".join(texts))}</p>'
    for p in d.paragraphs:
        parts.append(para_html(p))
    for t in d.tables:
        rows = []
        for r in t.rows:
            cells = []
            for c in r.cells:
                inner = ''.join(para_html(p) for p in c.paragraphs)
                cells.append(f'<td>{inner}</td>')
            rows.append('<tr>' + ''.join(cells) + '</tr>')
        parts.append('<table>' + ''.join(rows) + '</table>')
    return parse_page('<html><body>' + ''.join(parts) + '</body></html>', url), dict(where={}), {}


def parse_xlsx(data, url):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            cells = []
            for c in row:
                v = '' if c.value is None else str(c.value)
                out, r, _, _ = recover_run(v, c.font.name if c.font else None)
                attrs = f' data-jaani-rec-text="{r}"' if r else ''
                cells.append(f'<td{attrs}>{htmlmod.escape(out)}</td>')
            rows.append('<tr>' + ''.join(cells) + '</tr>')
        parts.append('<table>' + ''.join(rows) + '</table>')
    return parse_page('<html><body>' + ''.join(parts) + '</body></html>', url), dict(where={}), {}


def parse_file(data, url, ctype=''):
    low = (ctype or '').lower() + ' ' + url.lower()
    if 'pdf' in low:
        return parse_pdf(data, url)
    if 'wordprocessingml' in low or url.lower().endswith('.docx'):
        return parse_docx(data, url)
    if 'spreadsheetml' in low or url.lower().endswith('.xlsx'):
        return parse_xlsx(data, url)
    return None


FILE_LINK_RE = re.compile(r'\.(pdf|docx|xlsx)(?:$|\?)', re.I)
